"""Read-only structural inventory of a local data directory.

The module deliberately reports metadata and table headers, never raw row values. It contains no
OAI-specific filename, variable, visit, or linkage assumptions.
"""

from __future__ import annotations

import argparse
import csv
import json
from collections.abc import Callable, Sequence
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any


@dataclass(frozen=True, slots=True)
class TableMetadata:
    """Structural metadata returned by a supported tabular reader."""

    row_count: int
    column_names: tuple[str, ...]
    details: str | None = None


@dataclass(frozen=True, slots=True)
class FileInventory:
    """Safe-to-display structural metadata for one discovered file."""

    path: str
    extension: str
    size_bytes: int
    status: str
    row_count: int | None = None
    column_count: int | None = None
    column_names: tuple[str, ...] = ()
    details: str | None = None

    def as_dict(self) -> dict[str, Any]:
        """Return a JSON-serializable representation."""

        record = asdict(self)
        record["column_names"] = list(self.column_names)
        return record


def _header_names(values: Sequence[object]) -> tuple[str, ...]:
    """Convert only header cells to displayable names while preserving blank headers."""

    return tuple("" if value is None else str(value) for value in values)


def _inspect_delimited(path: Path, delimiter: str) -> TableMetadata:
    # csv.reader streams records, so even a large file is not loaded fully into memory.
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        try:
            header = next(reader)
        except StopIteration:
            return TableMetadata(row_count=0, column_names=())

        row_count = sum(1 for _ in reader)

    return TableMetadata(
        row_count=row_count,
        column_names=_header_names(header),
        details="First record treated as the header; decoded as UTF-8 with an optional BOM.",
    )


def _inspect_csv(path: Path) -> TableMetadata:
    return _inspect_delimited(path, delimiter=",")


def _inspect_tsv(path: Path) -> TableMetadata:
    return _inspect_delimited(path, delimiter="\t")


def _inspect_parquet(path: Path) -> TableMetadata:
    import pyarrow.parquet as parquet

    parquet_file = parquet.ParquetFile(path)
    return TableMetadata(
        row_count=parquet_file.metadata.num_rows,
        column_names=tuple(parquet_file.schema_arrow.names),
        details="Counts and schema read from Parquet metadata.",
    )


def _inspect_xlsx(path: Path) -> TableMetadata:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            header = ()
            row_count = 0
        else:
            # Read-only iteration avoids materializing the workbook; empty trailing rows are not
            # counted as observations.
            row_count = sum(1 for row in rows if any(value is not None for value in row))

        sheet_detail = f"Profiled worksheet {worksheet.title!r}."
        additional_sheet_count = len(workbook.sheetnames) - 1
        if additional_sheet_count:
            sheet_detail += (
                f" {additional_sheet_count} additional worksheet(s) were not structurally profiled."
            )

        return TableMetadata(
            row_count=row_count,
            column_names=_header_names(header),
            details=sheet_detail,
        )
    finally:
        workbook.close()


_TABLE_READERS: dict[str, Callable[[Path], TableMetadata]] = {
    ".csv": _inspect_csv,
    ".parquet": _inspect_parquet,
    ".tsv": _inspect_tsv,
    ".xlsx": _inspect_xlsx,
}


def _display_extension(path: Path) -> str:
    return "".join(path.suffixes).lower() or "(none)"


def _inspect_file(path: Path, root: Path) -> FileInventory:
    relative_path = path.relative_to(root).as_posix()
    extension = _display_extension(path)

    try:
        size_bytes = path.lstat().st_size
    except OSError as exc:
        return FileInventory(
            path=relative_path,
            extension=extension,
            size_bytes=0,
            status="error",
            details=f"Could not read file metadata: {exc}",
        )

    if path.is_symlink():
        return FileInventory(
            path=relative_path,
            extension=extension,
            size_bytes=size_bytes,
            status="skipped",
            details="Symbolic links are listed but not followed or inspected.",
        )

    reader = _TABLE_READERS.get(path.suffix.lower())
    if reader is None:
        return FileInventory(
            path=relative_path,
            extension=extension,
            size_bytes=size_bytes,
            status="unsupported",
            details=f"No structural reader configured for {extension}.",
        )

    try:
        table = reader(path)
    except Exception as exc:  # noqa: BLE001 - one corrupt file must not stop a directory inventory
        return FileInventory(
            path=relative_path,
            extension=extension,
            size_bytes=size_bytes,
            status="error",
            details=f"Structural inspection failed: {type(exc).__name__}: {exc}",
        )

    return FileInventory(
        path=relative_path,
        extension=extension,
        size_bytes=size_bytes,
        status="tabular",
        row_count=table.row_count,
        column_count=len(table.column_names),
        column_names=table.column_names,
        details=table.details,
    )


def inventory_directory(directory: str | Path) -> list[FileInventory]:
    """Recursively inventory files without changing or printing their contents.

    Paths in returned records are relative to ``directory``. Unsupported formats and symbolic
    links are reported rather than opened. A nonexistent path or a path that is not a directory
    raises a clear exception.
    """

    requested_root = Path(directory).expanduser()
    if not requested_root.exists():
        raise FileNotFoundError(f"Inventory directory does not exist: {requested_root}")
    if not requested_root.is_dir():
        raise NotADirectoryError(f"Inventory path is not a directory: {requested_root}")

    root = requested_root.resolve()
    files = sorted(
        (path for path in root.rglob("*") if path.is_file() or path.is_symlink()),
        key=lambda path: path.relative_to(root).as_posix(),
    )
    return [_inspect_file(path, root) for path in files]


def _human_size(size_bytes: int) -> str:
    size = float(size_bytes)
    for unit in ("B", "KiB", "MiB", "GiB", "TiB"):
        if size < 1024 or unit == "TiB":
            return f"{size:.0f} {unit}" if unit == "B" else f"{size:.1f} {unit}"
        size /= 1024
    raise AssertionError("unreachable")


def format_text_report(root: Path, records: Sequence[FileInventory]) -> str:
    """Format metadata only; raw table values are never accepted by this function."""

    lines = [f"Inventory root: {root.resolve()}", f"Files discovered: {len(records)}"]
    for record in records:
        summary = (
            f"[{record.status}] {record.path} | extension={record.extension} | "
            f"size={_human_size(record.size_bytes)}"
        )
        if record.status == "tabular":
            columns = ", ".join(record.column_names)
            summary += (
                f" | rows={record.row_count} | columns={record.column_count} "
                f"| column_names=[{columns}]"
            )
        if record.details:
            summary += f" | {record.details}"
        lines.append(summary)
    return "\n".join(lines)


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Recursively report file metadata and supported table structure without printing "
            "raw row values."
        )
    )
    parser.add_argument("directory", type=Path, help="Local directory to inspect recursively")
    parser.add_argument(
        "--json",
        action="store_true",
        dest="as_json",
        help="Emit machine-readable JSON instead of the text report",
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    """Command-line entry point."""

    parser = _build_parser()
    args = parser.parse_args(argv)
    try:
        records = inventory_directory(args.directory)
    except (FileNotFoundError, NotADirectoryError, OSError) as exc:
        parser.error(str(exc))

    if args.as_json:
        payload = {
            "root": str(args.directory.expanduser().resolve()),
            "files": [record.as_dict() for record in records],
        }
        print(json.dumps(payload, indent=2, sort_keys=True))
    else:
        print(format_text_report(args.directory, records))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
